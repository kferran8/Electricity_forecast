import numpy as np
from sklearn.linear_model import LinearRegression
from openpyxl import load_workbook

#ElectricityForecast

class OpenFileExel(object):
    def __init__(self):
        """Constructor"""
        self.filelink = None # ссылка файла
        self.worksheets = None # вызвает страницу
        self.page = 0
        self.max_column = None # максимальное количество столбцов данных
        self.max_row = None # максимальное количество строк данных
        self.data = None


    def open_file(self, filelink, page):
        self.filelink = filelink
        wb = load_workbook(filelink)
        self.worksheets = wb.worksheets[page]
        self.max_column = self.worksheets.max_column
        self.max_row = self.worksheets.max_row
        self.data = np.zeros((self.max_column, self.max_row))


class StatisticsData(object):
        pass



